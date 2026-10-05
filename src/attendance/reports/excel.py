"""Excel workbook for the monthly report."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..domain.worktime import DayRecord, EmployeeMonth


def monthly_xlsx(people: list[EmployeeMonth], title: str) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _write_table(
        summary,
        ["Employee No", "Name", "Days present", "Incomplete days", "Total hours", "Avg start", "Avg end"],
        [
            [p.employee_no, p.name, p.days_present, p.incomplete_days,
             round(p.total.total_seconds() / 3600, 2), p.avg_start, p.avg_end]
            for p in people
        ],
        {5: "0.00", 6: "HH:MM", 7: "HH:MM"},
    )

    daily = wb.create_sheet("Daily")
    records: list[DayRecord] = sorted((d for p in people for d in p.days), key=lambda d: (d.day, d.name))
    _write_table(
        daily,
        ["Date", "Employee No", "Name", "First scan", "Last scan", "Hours", "Scans", "Note"],
        [
            [d.day, d.employee_no, d.name, d.first.replace(tzinfo=None).time(),
             d.last.replace(tzinfo=None).time() if d.complete else None,
             round(d.worked.total_seconds() / 3600, 2), d.scans, "" if d.complete else "single scan"]
            for d in records
        ],
        {1: "YYYY-MM-DD", 4: "HH:MM", 5: "HH:MM", 6: "0.00"},
    )
    wb.properties.title = f"Attendance {title}"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _write_table(ws, headers: list[str], rows: list[list], formats: dict[int, str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for col, fmt in formats.items():
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = fmt
    for i, header in enumerate(headers, start=1):
        width = max([len(header)] + [len(str(r[i - 1] or "")) for r in rows])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 40)
    ws.freeze_panes = "A2"
